import hashlib
import json
import os
import shutil
import tempfile
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from dotenv import load_dotenv

from fastapi import FastAPI, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field

from sqlalchemy import (
    Column,
    DateTime,
    ForeignKey,
    Integer,
    String,
    Text,
    create_engine,
)
from sqlalchemy.orm import declarative_base, sessionmaker

from pypdf import PdfReader
from docx import Document as DocxDocument

from app.chunker import chunk_text
from app.embeddings import generate_embedding
from app.ingestion import DocumentIngestion
from app.retriever import Retriever
from app.vector_store import VectorStore


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv()

DB_SERVER = os.getenv(
    "DB_SERVER",
    r"DESKTOP-5EHP310\SQLEXPRESS",
)

DB_NAME = os.getenv(
    "DB_NAME",
    "EnterpriseRAG",
)

DB_DRIVER = os.getenv(
    "DB_DRIVER",
    "ODBC Driver 18 for SQL Server",
)

DATABASE_URL = (
    "mssql+pyodbc://@"
    + DB_SERVER
    + "/"
    + DB_NAME
    + "?driver="
    + DB_DRIVER.replace(" ", "+")
    + "&trusted_connection=yes"
    + "&TrustServerCertificate=yes"
)


# ============================================================
# DATABASE
# ============================================================

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,
)

SessionLocal = sessionmaker(
    autocommit=False,
    autoflush=False,
    bind=engine,
)

Base = declarative_base()


# ============================================================
# DOCUMENT MODEL
# ============================================================

class Document(Base):

    __tablename__ = "documents"

    id = Column(
        Integer,
        primary_key=True,
        index=True,
    )

    filename = Column(
        String(500),
        nullable=False,
    )

    file_size = Column(
        Integer,
        nullable=False,
        default=0,
    )

    chunk_count = Column(
        Integer,
        nullable=False,
        default=0,
    )

    file_hash = Column(
        String(64),
        nullable=False,
    )

    file_type = Column(
        String(20),
        nullable=False,
    )

    created_at = Column(
        DateTime,
        nullable=False,
        default=datetime.utcnow,
    )


# ============================================================
# DOCUMENT CHUNK MODEL
# ============================================================

class DocumentChunk(Base):

    __tablename__ = "document_chunks"

    id = Column(
        Integer,
        primary_key=True,
        index=True,
    )

    document_id = Column(
        Integer,
        ForeignKey("documents.id"),
        nullable=False,
        index=True,
    )

    chunk_text = Column(
        Text,
        nullable=False,
    )

    chunk_index = Column(
        Integer,
        nullable=False,
    )

    metadata_json = Column(
        Text,
        nullable=True,
    )

    created_at = Column(
        DateTime,
        nullable=False,
        default=datetime.utcnow,
    )


# ============================================================
# CHAT MESSAGE MODEL
# ============================================================

class ChatMessage(Base):

    __tablename__ = "chat_messages"

    id = Column(
        Integer,
        primary_key=True,
        index=True,
    )

    session_id = Column(
        String(100),
        nullable=False,
        index=True,
    )

    role = Column(
        String(20),
        nullable=False,
    )

    content = Column(
        Text,
        nullable=False,
    )

    created_at = Column(
        DateTime,
        nullable=False,
        default=datetime.utcnow,
    )


# ============================================================
# CREATE NEW TABLES IF NEEDED
# ============================================================

Base.metadata.create_all(
    bind=engine
)


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI(
    title="Enterprise RAG API",
    version="1.0.0",
)


# ============================================================
# CORS
# ============================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:5174",
        "http://127.0.0.1:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# VECTOR STORE
# ============================================================

vector_store = VectorStore(
    dimension=768,
)

retriever = Retriever(
    vector_store=vector_store,
)

rag_pipeline = None


# ============================================================
# CHAT REQUEST
# ============================================================

class ChatRequest(BaseModel):

    session_id: str = Field(
        ...,
        min_length=1,
    )

    message: str = Field(
        ...,
        min_length=1,
    )

    top_k: int = Field(
        default=5,
        ge=1,
        le=20,
    )


# ============================================================
# STARTUP
# ============================================================

@app.on_event("startup")
async def startup_event():

    global rag_pipeline

    from app.rag_pipeline import RAGPipeline

    rag_pipeline = RAGPipeline(
        retriever=retriever,
    )

    print(
        f"Vector store loaded with "
        f"{vector_store.count()} vectors."
    )


# ============================================================
# HEALTH
# ============================================================

@app.get("/health")
async def health():

    return {
        "status": "ok",
        "vector_count": vector_store.count(),
    }


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_pages(
    file_path: str,
) -> List[Dict[str, Any]]:

    reader = PdfReader(
        file_path,
    )

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        text = page.extract_text() or ""

        text = text.strip()

        if not text:
            continue

        pages.append(
            {
                "text": text,
                "page_number": page_number,
            }
        )

    return pages


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(
    file_path: str,
) -> str:

    document = DocxDocument(
        file_path,
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


# ============================================================
# CHAT HISTORY
# ============================================================

def get_chat_history(
    session_id: str,
    limit: int = 10,
):

    db = SessionLocal()

    try:

        messages = (
            db.query(ChatMessage)
            .filter(
                ChatMessage.session_id
                == session_id
            )
            .order_by(
                ChatMessage.created_at.desc()
            )
            .limit(limit)
            .all()
        )

        messages.reverse()

        return [
            {
                "role": message.role,
                "content": message.content,
            }
            for message in messages
        ]

    finally:

        db.close()


# ============================================================
# SAVE CHAT MESSAGE
# ============================================================

def save_chat_message(
    session_id: str,
    role: str,
    content: str,
):

    db = SessionLocal()

    try:

        message = ChatMessage(
            session_id=session_id,
            role=role,
            content=content,
        )

        db.add(message)

        db.commit()

    finally:

        db.close()


# ============================================================
# LIST DOCUMENTS
# ============================================================

@app.get("/documents")
async def get_documents():

    db = SessionLocal()

    try:

        documents = (
            db.query(Document)
            .order_by(
                Document.created_at.desc()
            )
            .all()
        )

        return [
            {
                "id": document.id,
                "filename": document.filename,
                "file_size": document.file_size,
                "chunk_count": document.chunk_count,
                "file_hash": document.file_hash,
                "file_type": document.file_type,
                "created_at": (
                    document.created_at.isoformat()
                    if document.created_at
                    else None
                ),
            }
            for document in documents
        ]

    except Exception as exc:

        print(
            "DOCUMENT LIST ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "database_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    finally:

        db.close()


# ============================================================
# UPLOAD DOCUMENT
# ============================================================

@app.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
):

    if not file.filename:

        return JSONResponse(
            status_code=400,
            content={
                "error": "validation_error",
                "detail": "Filename is required.",
            },
        )

    filename = Path(
        file.filename
    ).name

    extension = (
        Path(filename)
        .suffix
        .lower()
    )

    if extension not in (
        ".pdf",
        ".docx",
    ):

        return JSONResponse(
            status_code=400,
            content={
                "error": "validation_error",
                "detail": (
                    "Only PDF and DOCX files "
                    "are supported."
                ),
            },
        )

    db = SessionLocal()

    temp_dir = None

    try:

        # ----------------------------------------------------
        # READ FILE
        # ----------------------------------------------------

        contents = await file.read()

        if not contents:

            return JSONResponse(
                status_code=400,
                content={
                    "error": "empty_file",
                    "detail": (
                        "The uploaded file is empty."
                    ),
                },
            )

        file_size = len(contents)

        # ----------------------------------------------------
        # SHA-256 FILE HASH
        # ----------------------------------------------------

        file_hash = hashlib.sha256(
            contents
        ).hexdigest()

        print(
            f"File hash: {file_hash}"
        )

        # ----------------------------------------------------
        # DUPLICATE CHECK - FILENAME
        # ----------------------------------------------------

        existing_by_filename = (
            db.query(Document)
            .filter(
                Document.filename
                == filename
            )
            .first()
        )

        if existing_by_filename:

            return JSONResponse(
                status_code=409,
                content={
                    "error": "duplicate_document",
                    "detail": (
                        "A document with this "
                        "filename has already "
                        "been uploaded."
                    ),
                    "existing_document_id": (
                        existing_by_filename.id
                    ),
                },
            )

        # ----------------------------------------------------
        # DUPLICATE CHECK - FILE HASH
        # ----------------------------------------------------

        existing_by_hash = (
            db.query(Document)
            .filter(
                Document.file_hash
                == file_hash
            )
            .first()
        )

        if existing_by_hash:

            return JSONResponse(
                status_code=409,
                content={
                    "error": "duplicate_document",
                    "detail": (
                        "This exact document has "
                        "already been uploaded."
                    ),
                    "existing_document_id": (
                        existing_by_hash.id
                    ),
                    "existing_filename": (
                        existing_by_hash.filename
                    ),
                },
            )

        # ----------------------------------------------------
        # TEMPORARY DIRECTORY
        # ----------------------------------------------------

        temp_dir = tempfile.mkdtemp()

        temp_path = os.path.join(
            temp_dir,
            filename,
        )

        with open(
            temp_path,
            "wb",
        ) as output_file:

            output_file.write(
                contents
            )

        # ----------------------------------------------------
        # EXTRACT DOCUMENT
        # ----------------------------------------------------

        if extension == ".pdf":

            pages = extract_pdf_pages(
                temp_path
            )

        else:

            text = extract_docx_text(
                temp_path
            )

            pages = [
                {
                    "text": text,
                    "page_number": None,
                }
            ]

        # ----------------------------------------------------
        # CHUNK DOCUMENT
        # ----------------------------------------------------

        chunks = []

        for page in pages:

            page_chunks = chunk_text(
                page["text"]
            )

            for text in page_chunks:

                text = text.strip()

                if not text:
                    continue

                chunks.append(
                    {
                        "text": text,
                        "page_number": page[
                            "page_number"
                        ],
                    }
                )

        if not chunks:

            return JSONResponse(
                status_code=400,
                content={
                    "error": "empty_document",
                    "detail": (
                        "No readable text was "
                        "found in the document."
                    ),
                },
            )

        # ----------------------------------------------------
        # CREATE DOCUMENT
        # ----------------------------------------------------

        document = Document(
            filename=filename,
            file_size=file_size,
            chunk_count=len(chunks),
            file_hash=file_hash,
            file_type=extension.lstrip("."),
        )

        db.add(document)

        db.flush()

        # ----------------------------------------------------
        # STORE CHUNKS
        # ----------------------------------------------------

        chunk_texts = []

        metadata_list = []

        for index, chunk in enumerate(
            chunks
        ):

            metadata = {
                "document_id": document.id,
                "source": filename,
                "page_number": chunk[
                    "page_number"
                ],
                "section": None,
                "chunk_id": index,
            }

            chunk_record = DocumentChunk(
                document_id=document.id,
                chunk_text=chunk["text"],
                chunk_index=index,
                metadata_json=json.dumps(
                    metadata
                ),
            )

            db.add(
                chunk_record
            )

            chunk_texts.append(
                chunk["text"]
            )

            metadata_list.append(
                metadata
            )

        db.commit()

        # ----------------------------------------------------
        # GENERATE EMBEDDINGS
        # ----------------------------------------------------

        ingestion = DocumentIngestion(
            vector_store=vector_store,
        )

        await ingestion.add_chunks(
            chunks=chunk_texts,
            metadata=metadata_list,
        )

        # ----------------------------------------------------
        # SUCCESS RESPONSE
        # ----------------------------------------------------

        return {
            "message": (
                "Document uploaded successfully."
            ),
            "document": {
                "id": document.id,
                "filename": filename,
                "file_size": file_size,
                "chunk_count": len(chunks),
                "file_hash": file_hash,
                "file_type": extension.lstrip("."),
            },
        }

    except Exception as exc:

        db.rollback()

        print(
            "UPLOAD ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "upload_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    finally:

        db.close()

        if temp_dir:

            shutil.rmtree(
                temp_dir,
                ignore_errors=True,
            )


# ============================================================
# DELETE DOCUMENT
# ============================================================

@app.delete(
    "/documents/{document_id}"
)
async def delete_document(
    document_id: int,
):

    db = SessionLocal()

    try:

        document = (
            db.query(Document)
            .filter(
                Document.id
                == document_id
            )
            .first()
        )

        if not document:

            return JSONResponse(
                status_code=404,
                content={
                    "error": "not_found",
                    "detail": (
                        "Document not found."
                    ),
                },
            )

        filename = document.filename

        deleted_vectors = (
            vector_store.delete_by_document_id(
                document_id
            )
        )

        (
            db.query(DocumentChunk)
            .filter(
                DocumentChunk.document_id
                == document_id
            )
            .delete(
                synchronize_session=False
            )
        )

        db.delete(document)

        db.commit()

        return {
            "message": (
                "Document deleted successfully."
            ),
            "document_id": document_id,
            "filename": filename,
            "deleted_vectors": deleted_vectors,
        }

    except Exception as exc:

        db.rollback()

        print(
            "DELETE ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "delete_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    finally:

        db.close()


# ============================================================
# CHAT
# ============================================================

@app.post("/chat")
async def chat(
    request: ChatRequest,
):

    session_id = (
        request.session_id.strip()
    )

    message = (
        request.message.strip()
    )

    if not session_id:

        return JSONResponse(
            status_code=400,
            content={
                "error": "validation_error",
                "detail": (
                    "Session ID cannot be empty."
                ),
            },
        )

    if not message:

        return JSONResponse(
            status_code=400,
            content={
                "error": "validation_error",
                "detail": (
                    "Message cannot be empty."
                ),
            },
        )

    history = get_chat_history(
        session_id=session_id,
        limit=10,
    )

    try:

        result = await rag_pipeline.answer(
            question=message,
            top_k=request.top_k,
            history=history,
        )

    except Exception as exc:

        print(
            "RAG PIPELINE ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "rag_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    try:

        save_chat_message(
            session_id=session_id,
            role="user",
            content=message,
        )

        save_chat_message(
            session_id=session_id,
            role="assistant",
            content=result["answer"],
        )

    except Exception as exc:

        print(
            "CHAT HISTORY ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "database_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    return {
        "session_id": session_id,
        "question": message,
        "answer": result["answer"],
        "citations": result["citations"],
    }


# ============================================================
# VECTOR STORE INFORMATION
# ============================================================

@app.get("/vector-store")
async def vector_store_info():

    return {
        "vector_count": vector_store.count(),
        "dimension": vector_store.dimension,
    }


# ============================================================
# SYNC EXISTING FAISS VECTORS -> SQL SERVER
# ============================================================

@app.post(
    "/sync-vector-store-to-database"
)
async def sync_vector_store_to_database():

    db = SessionLocal()

    try:

        database_chunk_count = (
            db.query(DocumentChunk).count()
        )

        vector_count = (
            vector_store.count()
        )

        # ----------------------------------------------------
        # ALREADY SYNCHRONIZED
        # ----------------------------------------------------

        if database_chunk_count > 0:

            return {
                "message": (
                    "Database chunks already exist."
                ),
                "database_chunk_count": (
                    database_chunk_count
                ),
                "vector_count": vector_count,
            }

        if vector_count == 0:

            return JSONResponse(
                status_code=400,
                content={
                    "error": "empty_vector_store",
                    "detail": (
                        "No vectors exist to synchronize."
                    ),
                },
            )

        inserted = 0

        skipped = 0

        # ----------------------------------------------------
        # COPY EXISTING FAISS DATA
        # ----------------------------------------------------

        for index in range(
            vector_count
        ):

            text = vector_store.texts[
                index
            ]

            metadata = vector_store.metadata[
                index
            ]

            if not isinstance(
                metadata,
                dict,
            ):

                metadata = {}

            document_id = metadata.get(
                "document_id"
            )

            # ------------------------------------------------
            # DOCUMENT ID VALIDATION
            # ------------------------------------------------

            if document_id is None:

                print(
                    f"Skipping vector {index}: "
                    "missing document_id."
                )

                skipped += 1

                continue

            try:

                document_id = int(
                    document_id
                )

            except (
                ValueError,
                TypeError,
            ):

                print(
                    f"Skipping vector {index}: "
                    "invalid document_id."
                )

                skipped += 1

                continue

            document = (
                db.query(Document)
                .filter(
                    Document.id
                    == document_id
                )
                .first()
            )

            if not document:

                print(
                    f"Skipping vector {index}: "
                    f"document {document_id} "
                    "does not exist."
                )

                skipped += 1

                continue

            # ------------------------------------------------
            # CHUNK INDEX
            # ------------------------------------------------

            chunk_index = metadata.get(
                "chunk_id",
                index,
            )

            try:

                chunk_index = int(
                    chunk_index
                )

            except (
                ValueError,
                TypeError,
            ):

                chunk_index = index

            # ------------------------------------------------
            # INSERT CHUNK
            # ------------------------------------------------

            chunk_record = DocumentChunk(
                document_id=document_id,
                chunk_text=str(text),
                chunk_index=chunk_index,
                metadata_json=json.dumps(
                    metadata
                ),
            )

            db.add(
                chunk_record
            )

            inserted += 1

        db.commit()

        final_count = (
            db.query(DocumentChunk)
            .count()
        )

        return {
            "message": (
                "Existing vector store synchronized "
                "to database."
            ),
            "vector_count": vector_count,
            "inserted_chunks": inserted,
            "skipped_vectors": skipped,
            "database_chunk_count": final_count,
        }

    except Exception as exc:

        db.rollback()

        print(
            "VECTOR STORE SYNC ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "sync_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    finally:

        db.close()


# ============================================================
# REBUILD FAISS FROM SQL SERVER
# ============================================================

@app.post(
    "/rebuild-vector-store"
)
async def rebuild_vector_store():

    db = SessionLocal()

    rebuild_dir = "vector_data_rebuilt"

    try:

        chunks = (
            db.query(DocumentChunk)
            .order_by(
                DocumentChunk.document_id,
                DocumentChunk.chunk_index,
            )
            .all()
        )

        if not chunks:

            return {
                "message": (
                    "No stored chunks found."
                ),
                "vector_count": 0,
            }

        # ----------------------------------------------------
        # REMOVE OLD TEMP STORE
        # ----------------------------------------------------

        shutil.rmtree(
            rebuild_dir,
            ignore_errors=True,
        )

        # ----------------------------------------------------
        # CREATE TEMP VECTOR STORE
        # ----------------------------------------------------

        rebuilt_store = VectorStore(
            dimension=768,
            storage_dir=rebuild_dir,
        )

        embeddings = []

        texts = []

        metadata = []

        # ----------------------------------------------------
        # GENERATE EMBEDDINGS
        # ----------------------------------------------------

        for chunk in chunks:

            embedding = (
                await generate_embedding(
                    chunk.chunk_text
                )
            )

            embeddings.append(
                embedding
            )

            texts.append(
                chunk.chunk_text
            )

            if chunk.metadata_json:

                try:

                    chunk_metadata = json.loads(
                        chunk.metadata_json
                    )

                except json.JSONDecodeError:

                    chunk_metadata = {
                        "document_id": (
                            chunk.document_id
                        ),
                        "chunk_id": (
                            chunk.chunk_index
                        ),
                    }

            else:

                chunk_metadata = {
                    "document_id": (
                        chunk.document_id
                    ),
                    "chunk_id": (
                        chunk.chunk_index
                    ),
                }

            metadata.append(
                chunk_metadata
            )

        rebuilt_store.add(
            embeddings=embeddings,
            texts=texts,
            metadata=metadata,
        )

        # ----------------------------------------------------
        # REPLACE ACTIVE VECTOR STORE
        # ----------------------------------------------------

        vector_store.index = (
            rebuilt_store.index
        )

        vector_store.texts = (
            rebuilt_store.texts
        )

        vector_store.metadata = (
            rebuilt_store.metadata
        )

        vector_store.save()

        shutil.rmtree(
            rebuild_dir,
            ignore_errors=True,
        )

        return {
            "message": (
                "Vector store rebuilt successfully."
            ),
            "database_chunk_count": len(
                chunks
            ),
            "vector_count": (
                vector_store.count()
            ),
            "dimension": (
                vector_store.dimension
            ),
        }

    except Exception as exc:

        print(
            "VECTOR REBUILD ERROR:"
        )

        traceback.print_exc()

        return JSONResponse(
            status_code=500,
            content={
                "error": "rebuild_error",
                "detail": str(exc),
                "type": type(exc).__name__,
            },
        )

    finally:

        db.close()

        if os.path.exists(
            rebuild_dir
        ):

            shutil.rmtree(
                rebuild_dir,
                ignore_errors=True,
            )


# ============================================================
# GLOBAL ERROR HANDLER
# ============================================================

@app.exception_handler(Exception)
async def global_exception_handler(
    request,
    exc,
):

    print(
        "UNHANDLED API ERROR:"
    )

    traceback.print_exc()

    return JSONResponse(
        status_code=500,
        content={
            "error": "internal_server_error",
            "detail": str(exc),
            "type": type(exc).__name__,
        },
    )